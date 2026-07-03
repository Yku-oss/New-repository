import requests
from bs4 import BeautifulSoup
import time
from docx import Document
from docx.shared import Inches

# 请求头
headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/91.0.4472.124 Safari/537.36'
}

# 创建 Word 文档
doc = Document()
doc.add_heading('豆瓣电影 Top250 详细信息', 0)

# 存储所有电影信息（用于后续写入）
movies_data = []

# 循环爬取 10 页
for start in range(0, 250, 25):
    url = f'https://movie.douban.com/top250?start={start}'
    response = requests.get(url, headers=headers)
    soup = BeautifulSoup(response.text, 'html.parser')
    
    items = soup.select('.grid_view li')
    
    for item in items:
        # 获取详情页链接
        link = item.select_one('a')['href']
        
        # 请求详情页
        detail_response = requests.get(link, headers=headers)
        detail_soup = BeautifulSoup(detail_response.text, 'html.parser')
        
        # 排名
        rank = item.select_one('.pic em').text
        
        # 中文片名
        title = item.select_one('.title').text
        
        # 评分
        rating = item.select_one('.rating_num').text
        
        # 评分人数
        rating_num = item.select_one('.rating_num').find_next('span').text.strip() if item.select_one('.rating_num') else ''
        
        # 从详情页获取详细信息
        # 导演、编剧、主演等都在 id='info' 的 div 中
        info_div = detail_soup.find('div', id='info')
        if info_div:
            info_text = info_div.get_text(separator='\n').strip()
        else:
            info_text = ''
        
        # 类型（在详情页的 .indent 或 .pl 中找）
        type_tag = detail_soup.find('span', class_='pl')
        if type_tag and '类型' in type_tag.text:
            movie_type = type_tag.next_sibling.strip() if type_tag.next_sibling else ''
        else:
            movie_type = ''
        
        # 剧情简介
        summary_div = detail_soup.find('div', class_='related-info')
        if summary_div:
            summary = summary_div.get_text(strip=True)
        else:
            summary = ''
        
        # 上映时间、片长等可以从 info_div 中提取
        release_date = ''
        duration = ''
        if info_div:
            spans = info_div.find_all('span', class_='pl')
            for span in spans:
                text = span.text
                if '上映日期' in text:
                    release_date = span.next_sibling.strip() if span.next_sibling else ''
                elif '片长' in text:
                    duration = span.next_sibling.strip() if span.next_sibling else ''
        
        movies_data.append({
            'rank': rank,
            'title': title,
            'rating': rating,
            'rating_num': rating_num,
            'link': link,
            'info': info_text,
            'type': movie_type,
            'release_date': release_date,
            'duration': duration,
            'summary': summary
        })
        
        # 实时打印进度
        print(f"已爬取：{rank} - {title}")
        
        # 礼貌延时
        time.sleep(1)

# 写入 Word 文档
for movie in movies_data:
    doc.add_heading(f"{movie['rank']}. {movie['title']}", level=1)
    doc.add_paragraph(f"评分：{movie['rating']} ({movie['rating_num']}人评价)")
    doc.add_paragraph(f"链接：{movie['link']}")
    doc.add_paragraph(f"类型：{movie['type']}")
    doc.add_paragraph(f"上映日期：{movie['release_date']}")
    doc.add_paragraph(f"片长：{movie['duration']}")
    doc.add_paragraph("详细信息：")
    doc.add_paragraph(movie['info'])
    doc.add_paragraph("剧情简介：")
    doc.add_paragraph(movie['summary'])
    doc.add_paragraph("-" * 50)

# 保存文档
doc.save('豆瓣电影Top250详细信息.docx')
print("\n爬取完成！已保存到 豆瓣电影Top250详细信息.docx")