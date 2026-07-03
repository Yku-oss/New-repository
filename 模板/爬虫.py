import requests
from bs4 import BeautifulSoup
import time
import random
from docx import Document  # 新增：Word操作库

# 请求头（必须加，否则豆瓣会拦截）
ua = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
}

# 请求网页（实验要求保留的基础函数）
def page_request(url, ua):
    response = requests.get(url=url, headers=ua)
    html = response.content.decode('utf-8')
    return html

# 解析主页（获取电影名、排名、链接，用你原来的CSS选择器，符合实验1的基础）
def parse_index(html):
    soup = BeautifulSoup(html, 'lxml')
    # 获取排名、名称、评分、详情链接
    positions = soup.select('#content > div > div.article > ol > li > div > div.pic > em')
    names = soup.select('#content > div > div.article > ol > li > div > div.info > div.hd > a > span:nth-child(1)')
    ratings = soup.select('#content > div > div.article > ol > li > div > div.info > div.bd > div > span.rating_num')
    hrefs = soup.select('#content > div > div.article > ol > li > div > div.info > div.hd > a')
    
    # 把数据打包成列表，方便后续处理
    movie_list = []
    for i in range(len(names)):
        movie = {
            'rank': positions[i].get_text(),
            'name': names[i].get_text(),
            'rating': ratings[i].get_text(),
            'href': hrefs[i]['href']
        }
        movie_list.append(movie)
    return movie_list

# 【核心】解析子页面（实验2要求：用find()/find_all()，提取导演、编剧等详情）
def parse_detail(html):
    soup = BeautifulSoup(html, 'lxml')
    
    # 1. 提取info区域（实验明确要求：div id="info"）
    info = soup.find('div', id='info')
    if not info:
        return {}
    
    # 2. 封装提取info文本的工具函数（处理兄弟节点，符合实验要求）
    def get_info(label):
        tag = info.find('span', string=label)
        if tag and tag.next_sibling:
            # 清洗换行、空格，保证格式干净
            return tag.next_sibling.strip().replace('\n', '').replace(' ', '')
        return '无'
    
    # 3. 提取实验要求的所有字段
    director = get_info('导演')
    writer = get_info('编剧')
    actor = get_info('主演')
    type_ = get_info('类型')
    release = get_info('上映日期')
    duration = get_info('片长')
    
    # 4. 提取评分人数（独立于info区域）
    vote_tag = soup.find('span', property='v:votes')
    vote = vote_tag.get_text() if vote_tag else '0'
    
    # 5. 提取剧情简介
    summary_tag = soup.find('span', property='v:summary')
    summary = summary_tag.get_text().strip() if summary_tag else '无简介'
    
    # 6. 打包所有详情
    detail = {
        'director': director,
        'writer': writer,
        'actor': actor,
        'type': type_,
        'release': release,
        'duration': duration,
        'vote': vote,
        'summary': summary
    }
    return detail

# 主函数：完整爬取+保存
def main():
    print("**********开始爬取豆瓣电影Top250详情**********")
    # 1. 创建Word文档（实验要求保存到本地Word）
    doc = Document()
    doc.add_heading('豆瓣电影Top250 详细信息', 0)  # 一级标题
    
    # 2. 爬取10页（每页25部，共250部）
    for page in range(10):
        start = page * 25
        url = f'https://movie.douban.com/top250?start={start}'
        print(f'正在爬取第{page+1}页...')
        
        # 3. 请求主页，解析电影列表
        index_html = page_request(url, ua)
        movie_list = parse_index(index_html)
        
        # 4. 遍历每部电影，爬取详情页
        for movie in movie_list:
            # 4.1 请求详情页
            detail_html = page_request(movie['href'], ua)
            # 4.2 解析详情（用实验要求的find()/find_all()）
            detail = parse_detail(detail_html)
            
            # 4.3 写入Word
            doc.add_heading(f"{movie['rank']}. {movie['name']}", level=2)  # 二级标题
            doc.add_paragraph(f"评分：{movie['rating']}")
            doc.add_paragraph(f"导演：{detail.get('director', '无')}")
            doc.add_paragraph(f"编剧：{detail.get('writer', '无')}")
            doc.add_paragraph(f"主演：{detail.get('actor', '无')}")
            doc.add_paragraph(f"类型：{detail.get('type', '无')}")
            doc.add_paragraph(f"上映时间：{detail.get('release', '无')}")
            doc.add_paragraph(f"片长：{detail.get('duration', '无')}")
            doc.add_paragraph(f"评分人数：{detail.get('vote', '0')}人")
            doc.add_paragraph(f"剧情简介：\n{detail.get('summary', '无简介')}")
            doc.add_paragraph('-'*80)  # 分隔线
            
            # 4.4 防封：延时1秒（必须加，否则会被豆瓣拦截）
            time.sleep(1)
    
    # 5. 保存Word到本地（和你的爬虫.py同目录）
    doc.save('豆瓣Top250电影详细信息.docx')
    print("**********爬取完成！数据已保存到Word文档**********")

if __name__ == '__main__':
    main()